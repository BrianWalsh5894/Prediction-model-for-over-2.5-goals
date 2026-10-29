"""
Ongoing results spreadsheet: reports/live_tracking.xlsx

Rebuilt from the database every time weekly_settle.py runs (or run it on
its own), so it is always up to date - never edit it by hand, edits are
overwritten. Sheets:

  Summary      totals per league group and overall (live Excel formulas
               reading the Bets sheet)
  Bets         every paper bet: price taken, model probability, EV, result,
               profit, running profit, Betfair closing price and
               closing-line value (CLV)
  Predictions  every prediction, with your recorded price, the market's
               probability, Betfair close and the result

Run from inside src/:
    python export_tracking_xlsx.py
"""

import os
import sys
import sqlite3
from datetime import datetime
import pandas as pd

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "config")))
import config  # noqa: E402
import league_groups  # noqa: E402

DB_PATH = os.path.join("..", "data", "over25.db")
XLSX_PATH = os.path.join("..", "reports", "live_tracking.xlsx")

FONT = "Arial"


def load(conn):
    # scores typed in by manual_settle.py (used until football-data publishes them)
    conn.execute("""CREATE TABLE IF NOT EXISTS manual_results (
                        match_id TEXT PRIMARY KEY, home_goals INTEGER NOT NULL, away_goals INTEGER NOT NULL,
                        entered_at TEXT NOT NULL)""")
    bets = pd.read_sql_query("""
        SELECT b.bet_id, f.match_date, f.kickoff_time, f.league_name, f.home_team, f.away_team,
               b.bet_side, b.odds_taken, p.predicted_prob_over_2_5 AS p_over, b.expected_value,
               b.stake, b.commission, b.result, b.profit_loss, b.bet_date AS checked_at,
               o.bfe_close_over_2_5 AS close_over, o.bfe_close_under_2_5 AS close_under,
               COALESCE(m.home_goals, mr.home_goals) AS home_goals,
               COALESCE(m.away_goals, mr.away_goals) AS away_goals
        FROM bets b
        JOIN live_fixtures f ON f.match_id = b.match_id
        LEFT JOIN predictions p ON p.prediction_id = b.prediction_id
        LEFT JOIN odds o ON o.match_id = b.match_id
        LEFT JOIN matches m ON m.match_id = b.match_id
        LEFT JOIN manual_results mr ON mr.match_id = b.match_id
        ORDER BY f.match_date, f.kickoff_time, b.bet_id
    """, conn)
    preds = pd.read_sql_query("""
        SELECT f.match_date, f.kickoff_time, f.league_name, f.home_team, f.away_team, f.status,
               p.model_version, p.predicted_prob_over_2_5 AS p_over,
               lp.over_odds, lp.under_odds, lp.checked_at,
               o.bfe_close_over_2_5 AS close_over, o.bfe_close_under_2_5 AS close_under,
               COALESCE(m.home_goals, mr.home_goals) AS home_goals,
               COALESCE(m.away_goals, mr.away_goals) AS away_goals,
               COALESCE(m.over_2_5, CASE WHEN mr.match_id IS NULL THEN NULL
                                         ELSE (mr.home_goals + mr.away_goals > 2) END) AS over_2_5
        FROM predictions p
        JOIN live_fixtures f ON f.match_id = p.match_id
        LEFT JOIN live_prices lp ON lp.match_id = p.match_id
        LEFT JOIN odds o ON o.match_id = p.match_id
        LEFT JOIN matches m ON m.match_id = p.match_id
        LEFT JOIN manual_results mr ON mr.match_id = p.match_id
        ORDER BY f.match_date, f.kickoff_time, f.home_team
    """, conn)
    return bets, preds


def build(path=XLSX_PATH):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule

    with sqlite3.connect(DB_PATH) as conn:
        bets, preds = load(conn)
    group_of = league_groups.group_of_league

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_bets = wb.create_sheet("Bets")
    ws_pred = wb.create_sheet("Predictions")

    base = Font(name=FONT, size=10)
    bold = Font(name=FONT, size=10, bold=True)
    head_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="1F3864")
    green = PatternFill("solid", start_color="E2EFDA")
    red = PatternFill("solid", start_color="FCE4D6")
    thin = Border(bottom=Side(style="thin", color="BFBFBF"))

    def header(ws, cols, widths):
        for i, (c, w) in enumerate(zip(cols, widths), start=1):
            cell = ws.cell(row=1, column=i, value=c)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"

    # ---------------- Bets ----------------
    bet_cols = ["Date", "KO", "Group", "League", "Home", "Away", "Side", "Odds taken", "Model P (side)",
                "EV", "Stake", "Result", "Score", "Profit (units)", "Running profit", "Betfair close (side)",
                "CLV", "Price checked"]
    header(ws_bets, bet_cols, [11, 6, 8, 15, 20, 20, 7, 10, 11, 8, 7, 8, 7, 11, 11, 12, 8, 16])
    for i, r in enumerate(bets.itertuples(index=False), start=2):
        p_side = r.p_over if r.bet_side == "over" else (1 - r.p_over if pd.notna(r.p_over) else None)
        close = r.close_over if r.bet_side == "over" else r.close_under
        score = f"{int(r.home_goals)}-{int(r.away_goals)}" if pd.notna(r.home_goals) else ""
        values = [
            datetime.strptime(r.match_date, "%Y-%m-%d").date(), r.kickoff_time or "", group_of(r.league_name),
            r.league_name, r.home_team, r.away_team, r.bet_side.upper(), r.odds_taken, p_side,
            r.expected_value, r.stake, r.result, score,
            None if r.result == "pending" else r.profit_loss,
            f"=SUM($N$2:N{i})",
            None if pd.isna(close) else float(close),
            f'=IF(P{i}="","",H{i}/P{i}-1)',
            r.checked_at,
        ]
        for j, v in enumerate(values, start=1):
            c = ws_bets.cell(row=i, column=j, value=v)
            c.font, c.border = base, thin
        ws_bets.cell(row=i, column=1).number_format = "dd/mm/yyyy"
        for col in (9, 10, 17):
            ws_bets.cell(row=i, column=col).number_format = "0.0%"
        for col in (8, 16):
            ws_bets.cell(row=i, column=col).number_format = "0.00"
        for col in (14, 15):
            ws_bets.cell(row=i, column=col).number_format = '+0.00;-0.00;0.00'
    n_bets = max(len(bets), 1)
    last = n_bets + 1
    ws_bets.conditional_formatting.add(f"L2:L{last}", CellIsRule(operator="equal", formula=['"win"'], fill=green))
    ws_bets.conditional_formatting.add(f"L2:L{last}", CellIsRule(operator="equal", formula=['"loss"'], fill=red))
    ws_bets.auto_filter.ref = f"A1:R{last}"

    # ---------------- Predictions ----------------
    pred_cols = ["Date", "KO", "Group", "League", "Home", "Away", "Model", "Model P(Over)", "Your Over odds",
                 "Your Under odds", "Market P(Over)", "Betfair close P(Over)", "Score", "Over 2.5?", "Status",
                 "Price checked"]
    header(ws_pred, pred_cols, [11, 6, 8, 15, 20, 20, 7, 11, 10, 10, 11, 12, 7, 9, 10, 16])
    for i, r in enumerate(preds.itertuples(index=False), start=2):
        close_p = None
        if pd.notna(r.close_over) and pd.notna(r.close_under):
            close_p = (1 / r.close_over) / (1 / r.close_over + 1 / r.close_under)
        values = [
            datetime.strptime(r.match_date, "%Y-%m-%d").date(), r.kickoff_time or "", group_of(r.league_name),
            r.league_name, r.home_team, r.away_team, r.model_version, r.p_over,
            None if pd.isna(r.over_odds) else r.over_odds, None if pd.isna(r.under_odds) else r.under_odds,
            f'=IF(OR(I{i}="",J{i}=""),"",(1/I{i})/(1/I{i}+1/J{i}))',
            close_p,
            f"{int(r.home_goals)}-{int(r.away_goals)}" if pd.notna(r.home_goals) else "",
            ("" if pd.isna(r.over_2_5) else ("Yes" if r.over_2_5 == 1 else "No")),
            r.status, r.checked_at or "",
        ]
        for j, v in enumerate(values, start=1):
            c = ws_pred.cell(row=i, column=j, value=v)
            c.font, c.border = base, thin
        ws_pred.cell(row=i, column=1).number_format = "dd/mm/yyyy"
        for col in (8, 11, 12):
            ws_pred.cell(row=i, column=col).number_format = "0.0%"
        for col in (9, 10):
            ws_pred.cell(row=i, column=col).number_format = "0.00"
    ws_pred.auto_filter.ref = f"A1:P{max(len(preds), 1) + 1}"

    # ---------------- Summary (formulas over the Bets sheet) ----------------
    ws_sum.column_dimensions["A"].width = 30
    for col in "BCD":
        ws_sum.column_dimensions[col].width = 16
    ws_sum["A1"] = "Over 2.5 model v3 - live paper tracking"
    ws_sum["A1"].font = Font(name=FONT, size=14, bold=True)
    ws_sum["A2"] = (f"Rebuilt automatically by weekly_settle.py - last update {datetime.now():%d/%m/%Y %H:%M}. "
                    f"Do not edit: changes are overwritten.")
    ws_sum["A2"].font = Font(name=FONT, size=9, italic=True, color="808080")
    ws_sum["A3"] = (f"Rules: flat {config.PAPER_STAKE:g} unit paper stakes, {config.COMMISSION_RATE:.0%} "
                    f"commission, bet when EV >= {config.MIN_EDGE:.0%}. No real money until 200+ settled bets "
                    f"beat Betfair close with ROI interval above zero.")
    ws_sum["A3"].font = Font(name=FONT, size=9, italic=True, color="808080")

    groups = list(config.LEAGUE_GROUPS)
    cols = groups + ["Total"]
    for j, g in enumerate(cols, start=2):
        c = ws_sum.cell(row=5, column=j, value=g if g == "Total" else f"{g}\n({', '.join(config.LEAGUE_GROUPS[g]['leagues'])})")
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_sum.row_dimensions[5].height = 42
    ws_sum.cell(row=5, column=1, value="Metric").font = head_font
    ws_sum.cell(row=5, column=1).fill = head_fill

    R = f"Bets!$C$2:$C${last}"       # group
    RES = f"Bets!$L$2:$L${last}"     # result
    PL = f"Bets!$N$2:$N${last}"      # profit
    ST = f"Bets!$K$2:$K${last}"      # stake
    OD = f"Bets!$H$2:$H${last}"      # odds
    EV = f"Bets!$J$2:$J${last}"      # EV
    CLV = f"Bets!$Q$2:$Q${last}"     # CLV

    def crit(g):
        return f'"{g}"' if g != "Total" else '"*"'

    metrics = [
        ("Bets placed", lambda g: f"=COUNTIFS({R},{crit(g)})", "0"),
        ("Settled (win + loss)", lambda g: f'=COUNTIFS({R},{crit(g)},{RES},"win")+COUNTIFS({R},{crit(g)},{RES},"loss")', "0"),
        ("Pending", lambda g: f'=COUNTIFS({R},{crit(g)},{RES},"pending")', "0"),
        ("Void", lambda g: f'=COUNTIFS({R},{crit(g)},{RES},"void")', "0"),
        ("Wins", lambda g: f'=COUNTIFS({R},{crit(g)},{RES},"win")', "0"),
        ("Strike rate", lambda g: None, "0.0%"),
        ("Staked on settled bets (units)", lambda g: f'=SUMIFS({ST},{R},{crit(g)},{RES},"win")+SUMIFS({ST},{R},{crit(g)},{RES},"loss")', "0.00"),
        ("Profit (units)", lambda g: f"=SUMIFS({PL},{R},{crit(g)})", "+0.00;-0.00;0.00"),
        ("ROI", lambda g: None, "+0.0%;-0.0%;0.0%"),
        ("Average odds taken", lambda g: f'=IFERROR(AVERAGEIFS({OD},{R},{crit(g)}),"")', "0.00"),
        ("Average expected value (model)", lambda g: f'=IFERROR(AVERAGEIFS({EV},{R},{crit(g)}),"")', "+0.0%;-0.0%"),
        ("Average closing-line value", lambda g: f'=IFERROR(AVERAGEIFS({CLV},{R},{crit(g)},{CLV},"<>"),"")', "+0.0%;-0.0%"),
        ("Progress to 200 settled bets", lambda g: None, "0%"),
    ]
    row_of = {}
    for i, (label, _, _) in enumerate(metrics, start=6):
        row_of[label] = i
    for i, (label, fn, fmt) in enumerate(metrics, start=6):
        ws_sum.cell(row=i, column=1, value=label).font = bold
        for j, g in enumerate(cols, start=2):
            col = get_column_letter(j)
            f = fn(g)
            if label == "Strike rate":
                f = f'=IFERROR({col}{row_of["Wins"]}/{col}{row_of["Settled (win + loss)"]},"")'
            elif label == "ROI":
                f = f'=IFERROR({col}{row_of["Profit (units)"]}/{col}{row_of["Staked on settled bets (units)"]},"")'
            elif label == "Progress to 200 settled bets":
                f = f'=MIN(1,{col}{row_of["Settled (win + loss)"]}/200)'
            c = ws_sum.cell(row=i, column=j, value=f)
            c.font, c.number_format, c.border = base, fmt, thin
            c.alignment = Alignment(horizontal="center")
    note_row = 6 + len(metrics) + 1
    ws_sum.cell(row=note_row, column=1, value=(
        "Closing-line value = odds taken / Betfair closing odds - 1. Consistently positive CLV is the early "
        "sign of a real edge; profit over a few dozen bets is mostly luck.")).font = Font(name=FONT, size=9, italic=True)

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path, len(bets), len(preds)


def main():
    try:
        path, nb, npred = build()
    except ImportError:
        print("openpyxl is not installed - run:  pip install openpyxl   then try again.")
        sys.exit(1)
    except PermissionError:
        print(f"Couldn't write {XLSX_PATH} - it's open in Excel. Close it and run: python export_tracking_xlsx.py")
        sys.exit(1)
    print(f"Wrote {path} ({nb} bets, {npred} predictions)")


if __name__ == "__main__":
    main()
